from datetime import date, timedelta
from io import BytesIO

import fitz
import pytest
from docx import Document

from app.pipeline.indexer import PersistentLexicalIndex, split_passages
from app.pipeline.parser import DocumentParseError, parse_document
from app.pipeline.retriever import retrieve_passages


def test_parse_markdown_preserves_text_and_filename():
    parsed = parse_document(b"# Returns\n\nEligible items may be returned within 30 days.", "policy.md")

    assert parsed["fileName"] == "policy.md"
    assert parsed["format"] == "markdown"
    assert parsed["pages"] == [{"page": None, "text": "# Returns\n\nEligible items may be returned within 30 days."}]


def test_parse_pdf_preserves_page_numbers():
    document = fitz.open()
    document.new_page().insert_text((72, 72), "Page one policy")
    document.new_page().insert_text((72, 72), "Page two policy")
    pdf_bytes = document.tobytes()
    document.close()

    parsed = parse_document(pdf_bytes, "policy.pdf", "application/pdf")

    assert [page["page"] for page in parsed["pages"]] == [1, 2]
    assert "Page one policy" in parsed["text"]
    assert "Page two policy" in parsed["text"]


def test_parse_docx_includes_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Returns are accepted within 30 days.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Owner"
    table.rows[0].cells[1].text = "Legal"
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_document(buffer.getvalue(), "policy.docx")

    assert "Returns are accepted within 30 days." in parsed["text"]
    assert "Owner | Legal" in parsed["text"]


def test_empty_or_malformed_documents_have_structured_errors():
    with pytest.raises(DocumentParseError, match="NO_EXTRACTABLE_TEXT"):
        parse_document(b"", "empty.txt")
    with pytest.raises(DocumentParseError, match="INVALID_DOCUMENT"):
        parse_document(b"not a pdf", "broken.pdf")


def test_passage_ids_are_stable_and_page_aware():
    document = {
        "pages": [
            {"page": 2, "text": "First policy paragraph.\n\nSecond policy paragraph."},
        ],
    }
    source = {"orgId": "org-a", "sourceId": "source-1", "version": "1.0"}

    first = split_passages(document, source)
    second = split_passages(document, source)

    assert first == second
    assert first[0]["page"] == 2
    assert first[0]["sourceId"] == "source-1"
    assert first[0]["passageId"] != "source-1"


def test_index_persists_idempotently_and_retrieves_only_active_org_sources(tmp_path):
    index = PersistentLexicalIndex(tmp_path)
    document = {"pages": [{"page": 1, "text": "Eligible returns are accepted within 30 days."}]}
    source = {
        "orgId": "org-a",
        "sourceId": "source-a",
        "fileName": "a.md",
        "version": "1.0",
        "status": "ready",
        "expiryDate": (date.today() + timedelta(days=10)).isoformat(),
    }

    index.ingest(document, source)
    index.ingest(document, source)
    reloaded = PersistentLexicalIndex(tmp_path)

    assert len(reloaded.passages_for_org("org-a")) == 1
    assert retrieve_passages("returns 30 days", "org-a", reloaded)
    assert retrieve_passages("returns 30 days", "org-b", reloaded) == []


def test_retriever_excludes_expired_and_superseded_sources(tmp_path):
    index = PersistentLexicalIndex(tmp_path)
    document = {"pages": [{"page": 1, "text": "The old policy says returns within 30 days."}]}
    for source_id, status, expiry in (
        ("expired", "ready", (date.today() - timedelta(days=1)).isoformat()),
        ("superseded", "superseded", None),
    ):
        index.ingest(document, {
            "orgId": "org-a",
            "sourceId": source_id,
            "fileName": f"{source_id}.md",
            "version": "1.0",
            "status": status,
            "expiryDate": expiry,
        })

    assert retrieve_passages("old policy returns", "org-a", index) == []
